import os
import datetime
from docx import Document
from PyPDF2 import PdfReader
import openpyxl

EXCLUDED_DIRS = ["C:\\Windows", "C:\\Program Files", "C:\\Program Files (x86)", "C:\\$Recycle.Bin", "C:\\System Volume Information"]
EXCLUDED_KEYWORDS = ['.exe', '.dll', '.venv', '__pycache__', 'node_modules']

SUPPORTED_EXTENSIONS = (
    '.txt', '.docx', '.pdf', '.py', '.java', '.cpp', '.js', '.html', '.css',
    '.xlsx', '.xls', '.db'
)

def read_txt(file_path):
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return ""

def read_docx(file_path):
    try:
        doc = Document(file_path)
        return "\n".join(para.text for para in doc.paragraphs)
    except Exception:
        return ""

def read_pdf(file_path):
    try:
        reader = PdfReader(file_path)
        return "\n".join([page.extract_text() or "" for page in reader.pages])
    except Exception:
        return ""

def read_excel(file_path):
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        content = ""
        for sheet in wb:
            for row in sheet.iter_rows(values_only=True):
                content += " ".join([str(cell) if cell else "" for cell in row]) + "\n"
        return content
    except Exception:
        return ""

def read_code(file_path):
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return ""

def parse_file(file_path):
    if file_path.endswith(".txt"):
        return read_txt(file_path)
    elif file_path.endswith(".docx"):
        return read_docx(file_path)
    elif file_path.endswith(".pdf"):
        return read_pdf(file_path)
    elif file_path.endswith(".xlsx") or file_path.endswith(".xls"):
        return read_excel(file_path)
    elif file_path.endswith(('.py', '.java', '.cpp', '.js', '.html', '.css')):
        return read_code(file_path)
    elif file_path.endswith(".db"):
        return f"Database file: {os.path.basename(file_path)}"
    else:
        return None

def should_skip(path):
    lowered = path.lower()
    return any(ex in lowered for ex in EXCLUDED_KEYWORDS) or any(lowered.startswith(ex.lower()) for ex in EXCLUDED_DIRS)

def scan_and_parse_documents(base_dir, is_excluded):
    parsed_documents = {}
    print(f"📁 Scanning folder: {base_dir}")

    for root, _, files in os.walk(base_dir):
        if should_skip(root):
            continue
        for file in files:
            path = os.path.join(root, file)
            if should_skip(path):
                continue
            if file.lower().endswith(SUPPORTED_EXTENSIONS):
                print(f"🔍 Parsing: {path}")
                text = parse_file(path)
                if text:
                    parsed_documents[path] = text

    print("✅ Scanning and parsing complete.")
    return parsed_documents
